"""CV Extractor - Extract information from CV/resume files."""

from __future__ import annotations

import logging
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO

import docx
import openpyxl
from flask import Flask, render_template, request, send_file
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

logger = logging.getLogger(__name__)


@dataclass
class AppConfig:
    upload_folder: Path = Path("uploads")
    max_file_size: int = 16 * 1024 * 1024
    allowed_extensions: frozenset = field(
        default_factory=lambda: frozenset({".pdf", ".docx"}),
    )
    output_filename: str = "output.xlsx"


config = AppConfig()


def create_app() -> Flask:
    app = Flask(__name__)
    app.config.from_mapping(
        UPLOAD_FOLDER=str(config.upload_folder),
        MAX_CONTENT_LENGTH=config.max_file_size,
        SECRET_KEY="dev",
    )
    return app


app = create_app()


@dataclass
class CvData:
    name: str = ""
    email: str = ""
    phone: str = ""


def _allowed_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in config.allowed_extensions


def _extract_text_from_pdf(file: IO[bytes]) -> str:
    pdf_reader = PdfReader(file)
    return "\n".join(
        page.extract_text() or "" for page in pdf_reader.pages
    )


def _extract_text_from_docx(file_path: Path) -> str:
    doc = docx.Document(str(file_path))
    return "\n".join(para.text for para in doc.paragraphs)


def _extract_name(text: str) -> str:
    match = re.search(r"\b[A-Z][a-z]+(?:\s+[A-Z][a-z]+)+\b", text)
    return match.group() if match else ""


def _extract_email(text: str) -> str:
    match = re.search(r"[\w.%+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    return match.group() if match else ""


def _extract_phone(text: str) -> str:
    match = re.search(r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b", text)
    return match.group() if match else ""


def extract_data(file_path: Path) -> CvData:
    suffix = file_path.suffix.lower()

    try:
        if suffix == ".pdf":
            with open(file_path, "rb") as f:
                text = _extract_text_from_pdf(f)
        elif suffix == ".docx":
            text = _extract_text_from_docx(file_path)
        else:
            logger.warning("Unsupported file type: %s", suffix)
            return CvData()
    except (FileNotFoundError, PermissionError) as exc:
        logger.error("Failed to read %s: %s", file_path, exc)
        return CvData()

    if not text.strip():
        logger.warning("No text extracted from %s", file_path)
        return CvData()

    return CvData(
        name=_extract_name(text),
        email=_extract_email(text),
        phone=_extract_phone(text),
    )


def _build_workbook(results: list[CvData]) -> openpyxl.Workbook:
    workbook = openpyxl.Workbook()
    worksheet: Worksheet | None = workbook.active
    if worksheet is None:
        worksheet = workbook.create_sheet("CV Data")
    else:
        worksheet.title = "CV Data"
    worksheet.append(["Name", "Email", "Phone"])
    for data in results:
        worksheet.append([data.name, data.email, data.phone])
    return workbook


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/extract", methods=["POST"])
def extract():
    files: list[FileStorage] = request.files.getlist("files")
    if not files or all(f.filename == "" for f in files):
        return render_template("index.html", error="No files uploaded"), 400

    results: list[CvData] = []

    for file in files:
        filename = file.filename
        if not filename or not _allowed_file(filename):
            logger.warning("Skipping unsupported file: %s", filename)
            continue

        safe_name = secure_filename(filename)
        temp_dir = Path(tempfile.mkdtemp())
        file_path = temp_dir / safe_name

        try:
            file.save(str(file_path))
            data = extract_data(file_path)
            results.append(data)
            logger.info("Extracted data from %s: %s", filename, data)
        except OSError as exc:
            logger.error("File error for %s: %s", filename, exc)
        finally:
            try:
                file_path.unlink(missing_ok=True)
                temp_dir.rmdir()
            except OSError:
                pass

    if not results:
        return render_template("index.html", error="No valid data extracted"), 400

    workbook = _build_workbook(results)
    output_path = Path(tempfile.gettempdir()) / config.output_filename
    workbook.save(str(output_path))

    return send_file(
        str(output_path),
        as_attachment=True,
        download_name=config.output_filename,
    )


def configure_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def main() -> None:
    configure_logging()
    config.upload_folder.mkdir(exist_ok=True)
    logger.info("Starting CV Extractor on http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=5000, debug=True)


if __name__ == "__main__":
    main()
